"""Geração de relatórios Excel (.xlsx) e PDF."""
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


class GeradorRelatorios:
    @staticmethod
    def _estilo_cabecalho():
        return PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

    @staticmethod
    def excel(casos, caminho):
        """Gera relatório Excel formatado com os casos de um ciclo."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Execucao"

        headers = ["Código", "Tarefa", "Módulo", "Rotina", "Descrição", "Tipo", "Prioridade",
                   "Status", "% Concluído", "Responsável", "Horas Reais", "Evidência", "Observações"]
        ws.append(headers)

        header_fill = GeradorRelatorios._estilo_cabecalho()
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")

        for c in casos:
            ws.append([
                c.get("codigo"), c.get("tarefa"), c.get(
                    "modulo"), c.get("rotina"),
                c.get("descricao"), c.get("tipo"), c.get(
                    "prioridade"), c.get("status"),
                c.get("percentual"), c.get(
                    "responsavel"), c.get("horas_reais"),
                c.get("evidencia"), c.get("observacoes"),
            ])

        # Larguras
        larguras = [12, 22, 12, 14, 40, 12, 10, 14, 12, 14, 30, 30, 30]
        for i, w in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

        # Congelar cabeçalho
        ws.freeze_panes = "A2"

        # Filtro automático
        ws.auto_filter.ref = ws.dimensions

        wb.save(caminho)
        return caminho

    @staticmethod
    def pdf(casos, kpis, ciclo_nome, caminho):
        """Gera PDF executivo (landscape) com resumo e tabela de execução."""
        doc = SimpleDocTemplate(caminho, pagesize=landscape(A4),
                                leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
        styles = getSampleStyleSheet()
        elementos = []

        elementos.append(Paragraph(
            f"Relatório de Homologação — {ciclo_nome}", styles["Title"]))
        elementos.append(Spacer(1, 12))

        # Bloco de KPIs
        kpi_texto = (
            f"Total de casos: <b>{kpis['total']}</b> | "
            f"Execução: <b>{kpis['perc_execucao']}%</b> | "
            f"Aprovação: <b>{kpis['perc_aprovacao']}%</b> | "
            f"Defeitos: <b>{kpis['defeitos']}</b> | "
            f"Progresso médio: <b>{kpis['perc_medio']}%</b>"
        )
        elementos.append(Paragraph(kpi_texto, styles["Normal"]))
        elementos.append(Spacer(1, 16))

        # Tabela
        headers = ["Código", "Tarefa", "Módulo",
                   "Descrição", "Status", "%", "Responsável"]
        dados = [headers]
        for c in casos:
            dados.append([
                c.get("codigo", ""), c.get("tarefa", ""), c.get("modulo", ""),
                c.get("descricao", ""), c.get("status", ""),
                str(c.get("percentual", 0)), c.get("responsavel", ""),
            ])

        tabela = Table(dados, repeatRows=1)
        tabela.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#EAF1F8")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        elementos.append(tabela)

        elementos.append(Spacer(1, 16))
        elementos.append(Paragraph(
            f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} — Gerenciador Testes Essencial",
            styles["Italic"]))

        doc.build(elementos)
        return caminho
