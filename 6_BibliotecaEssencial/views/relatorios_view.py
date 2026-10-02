# ============================================
# views/relatorios_view.py
# Tela de RELATÓRIOS do Acervo (Biblioteca Essencial).
#
# - Filtros COMBINÁVEIS: Gênero, Editora e Situação
# - Pré-visualização em tabela (mesmo padrão das outras telas)
# - Total de livros encontrados (atualiza a cada filtro)
# - Exportação em DOIS formatos:
#     1) EXCEL (.xlsx) -> biblioteca openpyxl
#     2) PDF (.pdf)    -> biblioteca reportlab
#
# Dependências (instalar uma única vez):
#     pip install openpyxl reportlab
#
# Os arquivos gerados ficam na pasta "relatorios", criada
# automaticamente na raiz do projeto.
# ============================================
import os
from datetime import datetime

import customtkinter as ctk
from tkinter import ttk, messagebox

from database import Database
from repositories.editora_repository import EditoraRepository
from repositories.genero_repository import GeneroRepository


class RelatoriosView(ctk.CTkFrame):
    # Situação igual ao ASP.NET: 1 = Não Lido, 2 = Em Leitura, 3 = Concluído
    SITUACOES = {"1": "Não Lido", "2": "Em Leitura", "3": "Concluído"}
    SITUACAO_CODIGO = {"Não Lido": 1, "Em Leitura": 2, "Concluído": 3}

    COLUNAS = ("codigo", "nome_livro", "autor", "editora",
               "genero", "situacao", "data_compra")
    CABECALHOS = ("Código", "Título", "Autor", "Editora",
                  "Gênero", "Situação", "Data Compra")
    LARGURAS = (60, 230, 160, 160, 120, 100, 110)

    def __init__(self, master):
        super().__init__(master, fg_color="transparent")
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(2, weight=1)

        self._ultimas_linhas = []   # guarda o resultado atual para exportar
        self._gen_map = {}          # nome -> código (filtro de gênero)
        self._ed_map = {}           # nome -> código (filtro de editora)

        # Título
        ctk.CTkLabel(self, text="Relatório do Acervo",
                     font=("Arial", 22, "bold")).grid(
            row=0, column=0, sticky="w", padx=15, pady=(15, 5))

        # ===== Barra de filtros =====
        filtros = ctk.CTkFrame(self, fg_color="transparent")
        filtros.grid(row=1, column=0, sticky="ew", padx=15)

        ctk.CTkLabel(filtros, text="Gênero:",
                     font=("Segoe UI", 12)).pack(side="left", padx=(0, 4))
        self.filtro_genero = ctk.CTkOptionMenu(
            filtros, values=["Todos"], width=150)
        self.filtro_genero.set("Todos")
        self.filtro_genero.pack(side="left", padx=4)

        ctk.CTkLabel(filtros, text="Editora:",
                     font=("Segoe UI", 12)).pack(side="left", padx=(8, 4))
        self.filtro_editora = ctk.CTkOptionMenu(
            filtros, values=["Todos"], width=160)
        self.filtro_editora.set("Todos")
        self.filtro_editora.pack(side="left", padx=4)

        ctk.CTkLabel(filtros, text="Situação:",
                     font=("Segoe UI", 12)).pack(side="left", padx=(8, 4))
        self.filtro_situacao = ctk.CTkOptionMenu(
            filtros, values=["Todas", "Não Lido", "Em Leitura", "Concluído"],
            width=130)
        self.filtro_situacao.set("Todas")
        self.filtro_situacao.pack(side="left", padx=4)

        ctk.CTkButton(filtros, text="🔍 Gerar Relatório", width=150,
                      command=self._gerar).pack(side="left", padx=10)
        ctk.CTkButton(filtros, text="🧹 Limpar Filtros", width=130,
                      command=self._limpar).pack(side="left", padx=4)

        # ===== Tabela de pré-visualização =====
        self._montar_tabela()

        # ===== Rodapé: total + exportação =====
        rodape = ctk.CTkFrame(self, fg_color="transparent")
        rodape.grid(row=3, column=0, sticky="ew", padx=15, pady=(0, 15))

        self.lbl_total = ctk.CTkLabel(rodape, text="Total: 0 livro(s)",
                                      font=("Segoe UI", 13, "bold"),
                                      text_color="#2563EB")
        self.lbl_total.pack(side="left")

        ctk.CTkButton(rodape, text="📊 Exportar Excel", width=150,
                      command=self._exportar_excel).pack(side="right", padx=5)
        ctk.CTkButton(rodape, text="📄 Exportar PDF", width=140,
                      command=self._exportar_pdf).pack(side="right", padx=5)

        # Carrega as opções dos filtros e gera o relatório inicial
        self._carregar_opcoes()
        self._gerar()

    # ---------- Tabela ----------
    def _montar_tabela(self):
        self.estilo = ttk.Style()
        if "clam" in self.estilo.theme_names():
            self.estilo.theme_use("clam")
        self.estilo.configure(
            "Treeview", background="#FFFFFF", fieldbackground="#FFFFFF",
            foreground="#1F2937", borderwidth=0, rowheight=32,
            font=("Segoe UI", 12))
        self.estilo.configure(
            "Treeview.Heading", background="#E2E8F0", foreground="#0F172A",
            borderwidth=0, font=("Segoe UI", 12, "bold"))
        self.estilo.map("Treeview",
                        background=[("selected", "#2563EB")],
                        foreground=[("selected", "#FFFFFF")])

        table_frame = ctk.CTkFrame(self)
        table_frame.grid(row=2, column=0, sticky="nsew",
                         padx=15, pady=(10, 10))
        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)

        self.tree = ttk.Treeview(
            table_frame, columns=self.COLUNAS, show="headings",
            selectmode="browse")
        for col, cab, larg in zip(self.COLUNAS, self.CABECALHOS, self.LARGURAS):
            self.tree.heading(col, text=cab)
            self.tree.column(col, width=larg, minwidth=50,
                             anchor="center" if col == "codigo" else "w")

        self.tree.tag_configure("linha_par", background="#F1F5F9")
        self.tree.tag_configure("linha_impar", background="#FFFFFF")

        vsb = ttk.Scrollbar(table_frame, orient="vertical",
                            command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")

    # ---------- Opções dos filtros ----------
    def _carregar_opcoes(self):
        try:
            generos = GeneroRepository.list_all()
            editoras = EditoraRepository.list_all()
        except Exception as e:
            messagebox.showerror("Erro", f"Falha ao carregar opções:\n{e}")
            generos, editoras = [], []

        self._gen_map = {str(g[1]): g[0] for g in generos}
        self._ed_map = {str(e[1]): e[0] for e in editoras}

        self.filtro_genero.configure(
            values=["Todos"] + list(self._gen_map.keys()))
        self.filtro_genero.set("Todos")
        self.filtro_editora.configure(
            values=["Todos"] + list(self._ed_map.keys()))
        self.filtro_editora.set("Todos")

    # ---------- Consulta dinâmica (100% parametrizada) ----------
    def _consultar(self):
        clausulas = ["ISNULL(b.deletado, '') = ''"]
        params = []

        gen = self.filtro_genero.get()
        if gen != "Todos" and gen in self._gen_map:
            clausulas.append("b.genero = ?")
            params.append(self._gen_map[gen])

        ed = self.filtro_editora.get()
        if ed != "Todos" and ed in self._ed_map:
            clausulas.append("b.editora = ?")
            params.append(self._ed_map[ed])

        sit = self.filtro_situacao.get()
        if sit != "Todas" and sit in self.SITUACAO_CODIGO:
            clausulas.append("b.situacao = ?")
            params.append(self.SITUACAO_CODIGO[sit])

        sql = (
            "SELECT b.codigo, b.nome_livro, a.nome_autor, e.editora, "
            "g.genero, b.situacao, b.data_compra "
            "FROM Biblioteca b "
            "LEFT JOIN Autores a ON b.autor = a.codigo "
            "LEFT JOIN Editoras e ON b.editora = e.codigo "
            "LEFT JOIN Generos g ON b.genero = g.codigo "
            f"WHERE {' AND '.join(clausulas)} "
            "ORDER BY b.nome_livro"
        )
        return Database.query(sql, tuple(params))

    # ---------- Ações ----------
    def _gerar(self):
        try:
            rows = self._consultar()
        except Exception as e:
            messagebox.showerror("Erro", f"Falha ao consultar o acervo:\n{e}")
            rows = []
        self._ultimas_linhas = rows
        self._preencher(rows)

    def _limpar(self):
        self.filtro_genero.set("Todos")
        self.filtro_editora.set("Todos")
        self.filtro_situacao.set("Todas")
        self._gerar()

    def _preencher(self, rows):
        for item in self.tree.get_children():
            self.tree.delete(item)
        for indice, row in enumerate(rows):
            situacao = self.SITUACOES.get(str(row[5]), "")
            values = [row[0], row[1], row[2], row[3], row[4],
                      situacao, self._fmt_data(row[6])]
            tag = "linha_par" if indice % 2 == 0 else "linha_impar"
            self.tree.insert("", "end", values=values, tags=(tag,))
        self.lbl_total.configure(text=f"Total: {len(rows)} livro(s)")

    def _linhas_formatadas(self):
        """Linhas prontas para exportação (situação/datas legíveis)."""
        linhas = []
        for row in self._ultimas_linhas:
            situacao = self.SITUACOES.get(str(row[5]), "")
            linhas.append([
                row[0], row[1], row[2] or "", row[3] or "", row[4] or "",
                situacao or "", self._fmt_data(row[6]),
            ])
        return linhas

    # ---------- Exportação ----------
    @staticmethod
    def _caminho_arquivo(extensao):
        """Caminho do arquivo gerado: <projeto>/relatorios/relatorio_....ext"""
        raiz = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        pasta = os.path.join(raiz, "relatorios")
        os.makedirs(pasta, exist_ok=True)
        nome = f"relatorio_acervo_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{extensao}"
        return os.path.join(pasta, nome)

    def _exportar_excel(self):
        if not self._ultimas_linhas:
            messagebox.showwarning(
                "Aviso", "Gere o relatório antes de exportar.")
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import (Alignment, Border, Font,
                                         PatternFill, Side)
        except ImportError:
            messagebox.showerror(
                "Dependência faltando",
                "Instale a biblioteca openpyxl:\n\n    pip install openpyxl")
            return

        linhas = self._linhas_formatadas()
        caminho = self._caminho_arquivo("xlsx")

        wb = Workbook()
        ws = wb.active
        ws.title = "Acervo"

        # Título mesclado
        titulo = f"Relatório do Acervo - {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws.merge_cells(start_row=1, start_column=1, end_row=1,
                       end_column=len(self.CABECALHOS))
        c_titulo = ws.cell(row=1, column=1, value=titulo)
        c_titulo.font = Font(bold=True, size=14, color="1F2937")
        c_titulo.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 24

        # Cabeçalho
        ws.append(list(self.CABECALHOS))
        fill_cab = PatternFill("solid", fgColor="2563EB")
        font_cab = Font(bold=True, color="FFFFFF")
        thin = Side(style="thin", color="CBD5E1")
        borda = Border(left=thin, right=thin, top=thin, bottom=thin)

        for cel in ws[2]:
            cel.fill = fill_cab
            cel.font = font_cab
            cel.alignment = Alignment(horizontal="center", vertical="center")
            cel.border = borda

        # Dados com zebra
        fill_par = PatternFill("solid", fgColor="F1F5F9")
        fill_impar = PatternFill("solid", fgColor="FFFFFF")
        for i, linha in enumerate(linhas):
            ws.append(linha)
            for cel in ws[ws.max_row]:
                cel.border = borda
                cel.fill = fill_par if i % 2 == 0 else fill_impar

        # Larguras das colunas
        for idx, largura in enumerate(self.LARGURAS, start=1):
            ws.column_dimensions[chr(64 + idx)].width = max(largura // 7, 10)

        wb.save(caminho)
        messagebox.showinfo(
            "Exportado", f"Relatório Excel salvo em:\n{caminho}")

    def _exportar_pdf(self):
        if not self._ultimas_linhas:
            messagebox.showwarning(
                "Aviso", "Gere o relatório antes de exportar.")
            return
        try:
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib import colors
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.platypus import (Paragraph, SimpleDocTemplate,
                                            Table, TableStyle)
        except ImportError:
            messagebox.showerror(
                "Dependência faltando",
                "Instale a biblioteca reportlab:\n\n    pip install reportlab")
            return

        linhas = self._linhas_formatadas()
        caminho = self._caminho_arquivo("pdf")

        doc = SimpleDocTemplate(
            caminho,
            pagesize=landscape(A4),
            leftMargin=25, rightMargin=25,
            topMargin=25, bottomMargin=25,
            title="Relatório do Acervo - Biblioteca Essencial",
        )

        estilo_titulo = ParagraphStyle(
            "Titulo", fontName="Helvetica-Bold", fontSize=16,
            textColor=colors.HexColor("#1F2937"), spaceAfter=4)
        estilo_sub = ParagraphStyle(
            "Sub", fontName="Helvetica", fontSize=9,
            textColor=colors.HexColor("#64748B"), spaceAfter=12)

        titulo = "Relatório do Acervo"
        sub = (f"Gerado em {datetime.now().strftime('%d/%m/%Y às %H:%M')} - "
               f"Total de {len(linhas)} livro(s)")

        tabela = [list(self.CABECALHOS)] + [list(l) for l in linhas]
        t = Table(tabela, repeatRows=1)

        estilo_tabela = TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("ALIGN", (0, 0), (0, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F1F5F9")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ])
        t.setStyle(estilo_tabela)

        doc.build([Paragraph(titulo, estilo_titulo),
                   Paragraph(sub, estilo_sub),
                   t])
        messagebox.showinfo("Exportado", f"Relatório PDF salvo em:\n{caminho}")

    # ---------- Utilitários ----------
    @staticmethod
    def _fmt_data(valor):
        if not valor:
            return ""
        if isinstance(valor, datetime):
            return valor.strftime("%d/%m/%Y")
        return str(valor)
